from copy import deepcopy
from io import BytesIO
from pathlib import Path
import math
import re

import pandas as pd
import plotly.express as px
import streamlit as st
from openpyxl.styles import Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

DEFAULT_WORKBOOK = "foto_2025_python_met_uitvragen.xlsx"
ALL_CLUSTERS = "Alle clusters"
ALL_VARIANTS = "Alle varianten"
LEVEL_OPTIONS = ["Associate degree", "Bachelor", "Master"]

LABOUR_COLUMNS = [
    "Werkloosheid (%)",
    "Aantal maanden werkloos voor eerste baan",
    "Vaste aanstelling bij huidige functie (%)",
    "Dienstverband huidige functie: zelfstandige/freelance (%)",
    "Huidige functie binnen eigen of verwante opleidingsrichting (%)",
    "Bruto uurinkomen huidige functie (euro)",
]
DROP_COLUMNS = ["Bronbestand uitvraag", "Vraag 2: Aantal inschrijvingen correct?"]
VARIANT_RENAMES = {
    "Onderwijsinstelling": "INSTELLINGSNAAM ACTUEEL",
    "Gemeentenaam": "GEMEENTENAAM",
    "ISAT": "OPLEIDINGSCODE ACTUEEL",
    "Naam studie": "OPLEIDINGSNAAM ACTUEEL",
}
SCHOOL_COORDS = {
    ("Amsterdamse Hogeschool voor de Kunsten", "Amsterdam"): (52.3676, 4.9041),
    ("ArtEZ", "Arnhem"): (51.9851, 5.8987),
    ("ArtEZ", "Enschede"): (52.2215, 6.8937),
    ("ArtEZ", "Zwolle"): (52.5168, 6.0830),
    ("Avans Hogeschool", "Breda"): (51.5719, 4.7683),
    ("Avans Hogeschool", "'s-Hertogenbosch"): (51.6978, 5.3037),
    ("Avans Hogeschool", "Den Bosch"): (51.6978, 5.3037),
    ("Codarts, Hogeschool voor de Kunsten", "Rotterdam"): (51.9225, 4.4792),
    ("Design Academy Eindhoven", "Eindhoven"): (51.4416, 5.4697),
    ("Fontys Hogeschool", "Tilburg"): (51.5555, 5.0913),
    ("Gerrit Rietveld Academie", "Amsterdam"): (52.3676, 4.9041),
    ("Hanzehogeschool Groningen", "Amsterdam"): (52.3676, 4.9041),
    ("Hanzehogeschool Groningen", "Groningen"): (53.2194, 6.5665),
    ("Hanzehogeschool Groningen", "Leeuwarden"): (53.2012, 5.7999),
    ("Hogeschool Inholland", "Haarlem"): (52.3874, 4.6462),
    ("Hogeschool Leiden", "Leiden"): (52.1601, 4.4970),
    ("Hogeschool Rotterdam", "Rotterdam"): (51.9225, 4.4792),
    ("Hogeschool der Kunsten Den Haag", "'s-Gravenhage"): (52.0705, 4.3007),
    ("Hogeschool der Kunsten Den Haag", "Den Haag"): (52.0705, 4.3007),
    ("Hogeschool voor de Kunsten Utrecht", "Utrecht"): (52.0907, 5.1214),
    ("NHL Stenden Hogeschool", "Leeuwarden"): (53.2012, 5.7999),
    ("Saxion Hogeschool", "Deventer"): (52.2661, 6.1552),
    ("Zuyd Hogeschool", "Maastricht"): (50.8514, 5.6910),
}


def clean_text(value):
    if pd.isna(value):
        return ""
    text = str(value).strip()
    return "" if text.lower() in {"nan", "none", "<na>"} else text


def normalize_text(value):
    return re.sub(r"\s+", " ", clean_text(value)).casefold()


def normalize_city(value):
    city = clean_text(value)
    aliases = {
        "den bosch": "'s-Hertogenbosch", "s-hertogenbosch": "'s-Hertogenbosch",
        "'s hertogenbosch": "'s-Hertogenbosch", "den haag": "'s-Gravenhage",
        "s-gravenhage": "'s-Gravenhage", "'s gravenhage": "'s-Gravenhage",
    }
    return aliases.get(city.casefold(), city)


def clean_code(value):
    return re.sub(r"\.0$", "", clean_text(value))


def format_general_value(value):
    if pd.isna(value) or not clean_text(value):
        return "Niet beschikbaar"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else f"{value:g}".replace(".", ",")
    return clean_text(value)


def format_student_count(value):
    if pd.isna(value) or clean_text(value) == "":
        return "Niet beschikbaar"
    try:
        number = float(value)
        if number.is_integer():
            return f"{int(number):,}".replace(",", ".")
        return f"{number:,.1f}".replace(",", "X").replace(".", ",").replace("X", ".")
    except (TypeError, ValueError):
        return clean_text(value)


def format_value(value, column):
    if pd.isna(value) or clean_text(value) == "":
        return "Niet beschikbaar"
    if column == "Bruto uurinkomen huidige functie (euro)":
        try:
            return f"EUR {float(value):.2f}".replace(".", ",")
        except (TypeError, ValueError):
            return clean_text(value)
    if column.endswith("(%)"):
        try:
            return f"{float(value):g}%"
        except (TypeError, ValueError):
            return clean_text(value)
    return format_general_value(value)


def safe_filename(value):
    return "".join(c if c.isalnum() or c in "._- " else "_" for c in value)


def unique_clean_values(series):
    values = {}
    for value in series:
        text = clean_text(value)
        if text:
            values.setdefault(normalize_text(text), text)
    return sorted(values.values(), key=str.casefold)


# Opleidingsniveau wordt uitsluitend uit het begin van de opleidingsnaam afgeleid.
def get_education_level(value):
    name = clean_text(value)
    if re.match(r"^ad\s+", name, flags=re.IGNORECASE):
        return "Associate degree"
    if re.match(r"^b\s+", name, flags=re.IGNORECASE):
        return "Bachelor"
    if re.match(r"^m\s+", name, flags=re.IGNORECASE):
        return "Master"
    return "Onbekend"


def filter_levels(df, selected_levels):
    if not selected_levels:
        return df.iloc[0:0].copy()
    return df.loc[df["_education_level"].isin(selected_levels)].copy()


def show_level_filter():
    st.sidebar.markdown("### Opleidingsniveau")
    defaults = {
        "Associate degree": False,
        "Bachelor": True,
        "Master": False,
        "Onbekend": False,
    }
    selected = []
    for level in LEVEL_OPTIONS:
        key = "level_" + re.sub(r"\W+", "_", level.casefold()).strip("_")
        if key not in st.session_state:
            st.session_state[key] = defaults[level]
        if st.sidebar.checkbox(level, key=key):
            selected.append(level)
    if not selected:
        st.sidebar.warning("Selecteer minimaal één opleidingsniveau.")
    return selected


def split_cluster_values(value):
    text = clean_text(value)
    if not text:
        return []
    result, seen = [], set()
    for part in re.split(r"[\n\r;|]+", text):
        part = re.sub(r"^[•\-–—]+\s*", "", part)
        part = re.sub(r"\s+", " ", part).strip()
        key = normalize_text(part)
        if key and key not in seen:
            seen.add(key)
            result.append(part)
    return result


def get_cluster_values(df):
    values = {}
    for column in ["Cluster", "Cluster 2"]:
        if column in df.columns:
            for cell in df[column]:
                for value in split_cluster_values(cell):
                    values.setdefault(normalize_text(value), value)
    return sorted(values.values(), key=str.casefold)


def row_in_cluster(row, selected_cluster):
    selected_key = normalize_text(selected_cluster)
    return any(
        normalize_text(value) == selected_key
        for column in ["Cluster", "Cluster 2"] if column in row.index
        for value in split_cluster_values(row.get(column, ""))
    )


def filter_cluster(df, selected_cluster):
    if selected_cluster in {ALL_CLUSTERS, ALL_VARIANTS}:
        return df.copy()
    if not selected_cluster:
        return df.iloc[0:0].copy()
    return df.loc[df.apply(lambda row: row_in_cluster(row, selected_cluster), axis=1)].copy()


def get_discipline_values(df):
    return unique_clean_values(df["Discipline"]) if "Discipline" in df.columns else []


def filter_discipline(df, selected_discipline):
    if not selected_discipline:
        return df.iloc[0:0].copy()
    return df.loc[df["Discipline"].map(normalize_text).eq(normalize_text(selected_discipline))].copy()


def add_coordinates(df):
    df = df.copy()
    def lookup(row):
        school = clean_text(row.get("INSTELLINGSNAAM ACTUEEL", ""))
        city = normalize_city(row.get("GEMEENTENAAM", ""))
        exact = SCHOOL_COORDS.get((school, city))
        if exact:
            return exact
        matches = list(dict.fromkeys(v for (s, _), v in SCHOOL_COORDS.items() if s == school))
        return matches[0] if len(matches) == 1 else (None, None)
    coordinates = df.apply(lookup, axis=1, result_type="expand")
    df["Latitude"], df["Longitude"] = coordinates[0], coordinates[1]
    return df


def apply_coordinate_offsets(df):
    df = df.copy()
    if df.empty:
        return df
    df["Map Latitude"] = df["Latitude"].astype(float)
    df["Map Longitude"] = df["Longitude"].astype(float)
    for _, indices in df.groupby(["Latitude", "Longitude"], dropna=False, sort=False).groups.items():
        indices = list(indices)
        institutions = sorted({clean_text(df.at[i, "INSTELLINGSNAAM ACTUEEL"]) for i in indices})
        if len(institutions) <= 1:
            continue
        positions = {name: p for p, name in enumerate(institutions)}
        radius = 0.018 if len(institutions) <= 4 else 0.025
        for i in indices:
            position = positions[clean_text(df.at[i, "INSTELLINGSNAAM ACTUEEL"])]
            angle = 2 * math.pi * position / len(institutions)
            df.at[i, "Map Latitude"] += radius * math.sin(angle)
            df.at[i, "Map Longitude"] += radius * math.cos(angle)
    return df


def prepare_dataframe_keys(df):
    df = df.copy()
    df["_institution_key"] = df["INSTELLINGSNAAM ACTUEEL"].map(normalize_text)
    df["_code_key"] = df["OPLEIDINGSCODE ACTUEEL"].map(clean_code).map(normalize_text)
    df["_study_key"] = df["OPLEIDINGSNAAM ACTUEEL"].map(normalize_text)
    df["_city_key"] = df["GEMEENTENAAM"].map(normalize_city).map(normalize_text)
    df["_education_level"] = df["OPLEIDINGSNAAM ACTUEEL"].map(get_education_level)
    return df


@st.cache_data(show_spinner=False)
def read_workbook(source):
    opleidingen = pd.read_excel(source, sheet_name="opleidingen", engine="openpyxl")
    varianten = pd.read_excel(source, sheet_name="varianten", engine="openpyxl")
    opleidingen.columns = opleidingen.columns.map(lambda x: str(x).strip())
    varianten.columns = varianten.columns.map(lambda x: str(x).strip())
    opleidingen = opleidingen.drop(columns=DROP_COLUMNS, errors="ignore")
    varianten = varianten.drop(columns=DROP_COLUMNS, errors="ignore").rename(columns=VARIANT_RENAMES)
    required_o = ["INSTELLINGSNAAM ACTUEEL", "OPLEIDINGSCODE ACTUEEL", "OPLEIDINGSNAAM ACTUEEL", "GEMEENTENAAM", "Discipline", "Cluster", "Cluster 2"]
    required_v = ["INSTELLINGSNAAM ACTUEEL", "OPLEIDINGSCODE ACTUEEL", "OPLEIDINGSNAAM ACTUEEL", "GEMEENTENAAM", "Variant", "Aantal inschrijvingen 2025", "Cluster", "Cluster 2"]
    missing_o = [c for c in required_o if c not in opleidingen.columns]
    missing_v = [c for c in required_v if c not in varianten.columns]
    if missing_o:
        raise ValueError("In opleidingen ontbreken: " + ", ".join(missing_o))
    if missing_v:
        raise ValueError("In varianten ontbreken: " + ", ".join(missing_v))
    for df in [opleidingen, varianten]:
        df["OPLEIDINGSCODE ACTUEEL"] = df["OPLEIDINGSCODE ACTUEEL"].map(clean_code).astype("string")
        df["INSTELLINGSNAAM ACTUEEL"] = df["INSTELLINGSNAAM ACTUEEL"].astype("string").str.strip()
        df["OPLEIDINGSNAAM ACTUEEL"] = df["OPLEIDINGSNAAM ACTUEEL"].astype("string").str.strip()
        df["GEMEENTENAAM"] = df["GEMEENTENAAM"].map(normalize_city)
    opleidingen = prepare_dataframe_keys(opleidingen)
    varianten = prepare_dataframe_keys(varianten)
    available = [c for c in LABOUR_COLUMNS if c in opleidingen.columns]
    lookup = opleidingen[["_institution_key", "_code_key", "_study_key", "_city_key", "GEMEENTENAAM", *available]].drop_duplicates()
    varianten = varianten.drop(columns=LABOUR_COLUMNS, errors="ignore")
    with_city = varianten[varianten["_city_key"].ne("")].copy()
    without_city = varianten[varianten["_city_key"].eq("")].copy()
    if not with_city.empty:
        with_city = with_city.merge(
            lookup[["_institution_key", "_code_key", "_study_key", "_city_key", *available]].drop_duplicates(),
            how="left", on=["_institution_key", "_code_key", "_study_key", "_city_key"])
    if not without_city.empty:
        without_city = without_city.drop(columns=["GEMEENTENAAM", "_city_key"], errors="ignore").merge(
            lookup[["_institution_key", "_code_key", "_study_key", "GEMEENTENAAM", "_city_key", *available]].drop_duplicates(),
            how="left", on=["_institution_key", "_code_key", "_study_key"])
    varianten = pd.concat([with_city, without_city], ignore_index=True, sort=False)
    return add_coordinates(opleidingen), add_coordinates(varianten)


def show_labour_panel(df):
    """Toon altijd het gemiddelde van de unieke studies die op de kaart staan."""
    st.sidebar.markdown("### Gemiddelde arbeidsmarktinformatie")
    st.sidebar.caption(
        "Gemiddelde over alle unieke studies die op de kaart zichtbaar zijn "
        "en waarvoor een waarde beschikbaar is."
    )

    # Een opleiding kan meerdere opleidingsvormen of varianten hebben. Dezelfde
    # opleiding telt daarom maar één keer mee in het arbeidsmarktgemiddelde.
    study_key_columns = [
        column
        for column in ["_institution_key", "_code_key", "_city_key"]
        if column in df.columns
    ]
    average_df = (
        df.drop_duplicates(subset=study_key_columns)
        if study_key_columns
        else df.drop_duplicates()
    )

    for column in LABOUR_COLUMNS:
        if column not in average_df.columns:
            text = "Niet beschikbaar"
        else:
            numeric_values = pd.to_numeric(
                average_df[column], errors="coerce"
            ).dropna()
            if numeric_values.empty:
                text = "Niet beschikbaar"
            else:
                average = numeric_values.mean()
                if column == "Bruto uurinkomen huidige functie (euro)":
                    text = f"EUR {average:.2f}".replace(".", ",")
                elif column.endswith("(%)"):
                    text = f"{average:.1f}%".replace(".", ",")
                else:
                    text = f"{average:.1f}".replace(".", ",")
        st.sidebar.metric(column, text)

def show_list(df, column, heading, empty_text):
    st.sidebar.markdown(f"### {heading}")
    values = unique_clean_values(df[column]) if not df.empty and column in df.columns else []
    if not values:
        st.sidebar.caption(empty_text)
    for value in values:
        st.sidebar.markdown(f"- {value}")


def selection_to_excel(opleidingen_df, varianten_df):
    """Exporteer opleidingen en varianten naar twee opgemaakte tabbladen."""
    output = BytesIO()

    def prepare_export(df):
        return df.drop(
            columns=[
                column
                for column in df.columns
                if column.startswith("_")
                or column in {
                    "Latitude",
                    "Longitude",
                    "Map Latitude",
                    "Map Longitude",
                }
            ],
            errors="ignore",
        ).copy()

    sheets = {
        "Opleidingen": prepare_export(opleidingen_df),
        "Varianten": prepare_export(varianten_df),
    }

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for sheet_name, export_df in sheets.items():
            export_df.to_excel(
                writer,
                sheet_name=sheet_name,
                index=False,
            )

            worksheet = writer.book[sheet_name]
            worksheet.freeze_panes = "A2"

            # Tekstterugloop voor alle cellen en verticale uitlijning bovenaan.
            for row in worksheet.iter_rows():
                for cell in row:
                    cell.alignment = Alignment(
                        wrap_text=True,
                        vertical="top",
                    )

            # Beperk kolombreedtes zodat lange tekst daadwerkelijk terugloopt.
            for column_number, cells in enumerate(
                worksheet.iter_cols(),
                start=1,
            ):
                cells = list(cells)
                maximum_length = max(
                    (
                        len(str(cell.value))
                        if cell.value is not None
                        else 0
                    )
                    for cell in cells[:200]
                )
                worksheet.column_dimensions[
                    get_column_letter(column_number)
                ].width = min(
                    max(maximum_length + 2, 12),
                    35,
                )

            # Maak van ieder gevuld tabblad een echte Excel-tabel.
            if not export_df.empty:
                table_reference = (
                    f"A1:{get_column_letter(worksheet.max_column)}"
                    f"{worksheet.max_row}"
                )
                table_name = (
                    "TabelOpleidingen"
                    if sheet_name == "Opleidingen"
                    else "TabelVarianten"
                )
                table = Table(
                    displayName=table_name,
                    ref=table_reference,
                )
                table.tableStyleInfo = TableStyleInfo(
                    name="TableStyleMedium2",
                    showFirstColumn=False,
                    showLastColumn=False,
                    showRowStripes=True,
                    showColumnStripes=False,
                )
                worksheet.add_table(table)

    return output.getvalue()

def create_selected_study_keys(df):
    code_keys, name_keys = set(), set()
    for _, row in df.iterrows():
        institution, code, study = row.get("_institution_key", ""), row.get("_code_key", ""), row.get("_study_key", "")
        if institution and code:
            code_keys.add((institution, code))
        if institution and study:
            name_keys.add((institution, study))
    return code_keys, name_keys


def variants_for_selected_studies(selected_opleidingen, varianten):
    """Vind alle varianten die onder de geselecteerde opleidingen vallen."""
    code_keys, name_keys = create_selected_study_keys(selected_opleidingen)

    def belongs(row):
        institution = row.get("_institution_key", "")
        code = row.get("_code_key", "")
        study = row.get("_study_key", "")
        return bool(
            (institution and code and (institution, code) in code_keys)
            or (institution and study and (institution, study) in name_keys)
        )

    return varianten.loc[varianten.apply(belongs, axis=1)].copy()


def studies_for_selected_variants(opleidingen, selected_variants):
    """Vind de opleidingen die bij de actuele variantselectie horen."""
    code_keys, name_keys = create_selected_study_keys(selected_variants)
    def belongs(row):
        institution = row.get("_institution_key", "")
        code = row.get("_code_key", "")
        study = row.get("_study_key", "")
        return bool(
            (institution and code and (institution, code) in code_keys)
            or (institution and study and (institution, study) in name_keys)
        )
    return opleidingen.loc[opleidingen.apply(belongs, axis=1)].copy()


def variants_for_education_cluster(opleidingen, varianten, selected_cluster):
    selected = filter_cluster(opleidingen, selected_cluster)
    code_keys, name_keys = create_selected_study_keys(selected)
    def belongs(row):
        institution, code, study = row.get("_institution_key", ""), row.get("_code_key", ""), row.get("_study_key", "")
        return bool((institution and code and (institution, code) in code_keys) or (institution and study and (institution, study) in name_keys))
    return selected, varianten.loc[varianten.apply(belongs, axis=1)].copy()


def build_hover(group, mode):
    institution = clean_text(group["INSTELLINGSNAAM ACTUEEL"].iloc[0])
    city = clean_text(group["GEMEENTENAAM"].iloc[0])
    lines = [f"<b>{institution}</b>"] + ([city] if city else []) + [""]
    sort_columns = [c for c in (["OPLEIDINGSNAAM ACTUEEL", "OPLEIDINGSVORM"] if mode == "opleidingen" else ["OPLEIDINGSNAAM ACTUEEL", "Variant"]) if c in group.columns]
    sorted_group = group.sort_values(sort_columns, kind="stable") if sort_columns else group
    for position, (_, row) in enumerate(sorted_group.iterrows(), 1):
        if mode == "opleidingen":
            name = clean_text(row.get("OPLEIDINGSNAAM ACTUEEL")) or "Onbekende opleiding"
            lines += [f"<b>{position}. {name}</b>",
                      "Opleidingsvorm: " + format_general_value(row.get("OPLEIDINGSVORM")),
                      "Aantal studenten 2023: " + format_student_count(row.get("2023")),
                      "Aantal studenten 2024: " + format_student_count(row.get("2024")),
                      "Aantal studenten 2025: " + format_student_count(row.get("2025")),
                      "3-jarig gemiddelde studenten: " + format_student_count(row.get("3 jarig gemiddelde studenten")),
                      "Varianten studie: " + format_general_value(row.get("Varianten"))]
        else:
            variant = clean_text(row.get("Variant")) or "Onbekende variant"
            study = clean_text(row.get("OPLEIDINGSNAAM ACTUEEL")) or "Onbekende studie"
            lines += [f"<b>{position}. {variant}</b>", f"Naam studie: {study}",
                      "Aantal inschrijvingen 2025: " + format_student_count(row.get("Aantal inschrijvingen 2025"))]
        if position < len(sorted_group):
            lines.append("")
    return "<br>".join(lines)


def aggregate_map_data(df, mode):
    mappable = df.dropna(subset=["Latitude", "Longitude"]).copy()
    if mappable.empty:
        return pd.DataFrame()
    records = []
    group_columns = ["INSTELLINGSNAAM ACTUEEL", "GEMEENTENAAM", "Latitude", "Longitude"]
    for (institution, city, latitude, longitude), group in mappable.groupby(group_columns, dropna=False, sort=True):
        candidates = (["INSTELLINGSNAAM ACTUEEL", "GEMEENTENAAM", "OPLEIDINGSCODE ACTUEEL", "OPLEIDINGSNAAM ACTUEEL", "OPLEIDINGSVORM", "2023", "2024", "2025", "3 jarig gemiddelde studenten", "Varianten"]
                      if mode == "opleidingen" else
                      ["INSTELLINGSNAAM ACTUEEL", "GEMEENTENAAM", "OPLEIDINGSCODE ACTUEEL", "OPLEIDINGSNAAM ACTUEEL", "Variant", "Aantal inschrijvingen 2025"])
        display = group.drop_duplicates(subset=[c for c in candidates if c in group.columns])
        count = len(display)
        singular = "opleiding" if mode == "opleidingen" else "variant"
        plural = "opleidingen" if mode == "opleidingen" else "varianten"
        source_label = "Opleiding" if mode == "opleidingen" else "Variant"
        hover_text = f"<b>Type: {source_label}</b><br>" + build_hover(display, mode)
        records.append({"INSTELLINGSNAAM ACTUEEL": clean_text(institution), "GEMEENTENAAM": clean_text(city),
                        "Latitude": float(latitude), "Longitude": float(longitude), "Aantal onderdelen": count,
                        "Aantal label": f"{count} {singular if count == 1 else plural}",
                        "Marker label": str(count) if count > 1 else "", "Hover": hover_text})
    return apply_coordinate_offsets(pd.DataFrame(records))


def create_map_figure(map_data, title):
    institutions = sorted(unique_clean_values(map_data["INSTELLINGSNAAM ACTUEEL"]), key=str.casefold)
    max_count = max(int(map_data["Aantal onderdelen"].max()), 1)
    fig = px.scatter_map(map_data, lat="Map Latitude", lon="Map Longitude", color="INSTELLINGSNAAM ACTUEEL",
                         size="Aantal onderdelen", size_max=30 if max_count > 1 else 22,
                         category_orders={"INSTELLINGSNAAM ACTUEEL": institutions}, hover_name=None,
                         hover_data={"Map Latitude": False, "Map Longitude": False, "INSTELLINGSNAAM ACTUEEL": False,
                                     "Aantal onderdelen": False, "GEMEENTENAAM": False, "Hover": False},
                         custom_data=["Hover", "Marker label", "Aantal onderdelen"], height=760, title=title)
    original_traces = list(fig.data)
    fig.data = ()
    for trace in original_traces:
        outline = deepcopy(trace)
        outline.name, outline.showlegend, outline.hoverinfo, outline.hovertemplate = f"{trace.name} outline", False, "skip", None
        try:
            sizes = [float(size) + 7 for size in trace.marker.size]
        except TypeError:
            sizes = float(trace.marker.size or 19) + 7
        outline.marker = {"size": sizes, "color": "black", "opacity": 0.95}
        trace.marker.opacity = 0.90
        trace.hovertemplate = "%{customdata[0]}<extra></extra>"
        fig.add_trace(outline)
        fig.add_trace(trace)
    multiple = map_data[map_data["Aantal onderdelen"] > 1]
    if not multiple.empty:
        fig.add_scattermap(lat=multiple["Map Latitude"], lon=multiple["Map Longitude"], mode="text",
                           text=multiple["Marker label"], textfont={"size": 11, "color": "white", "family": "Arial Black"},
                           hoverinfo="skip", showlegend=False, name="Aantal onderdelen")
    fig.update_layout(paper_bgcolor="#eeeeee", plot_bgcolor="#eeeeee",
                      map={"style": "carto-positron", "center": {"lat": 52.20, "lon": 5.30}, "zoom": 5.35, "pitch": 0},
                      margin={"r": 0, "t": 55, "l": 0, "b": 0}, legend_title_text="Onderwijsinstelling",
                      hoverlabel={"bgcolor": "white", "font_size": 13, "font_family": "Arial", "align": "left"})
    return fig


def choose_discipline(level_filtered):
    disciplines = get_discipline_values(level_filtered)
    if not disciplines:
        st.warning("Binnen de gekozen opleidingsniveaus zijn geen disciplines gevonden.")
        st.stop()
    if st.session_state.get("selected_discipline") not in disciplines:
        st.session_state.selected_discipline = disciplines[0]
    st.sidebar.markdown("### Discipline")
    selected = st.sidebar.selectbox("Discipline", disciplines, key="selected_discipline", label_visibility="collapsed")
    return selected, filter_discipline(level_filtered, selected)


def combine_map_sources(studies_df, variants_df):
    """Combineer opleidingen en varianten tot één marker per instelling en locatie."""
    study_map = aggregate_map_data(studies_df, "opleidingen") if studies_df is not None and not studies_df.empty else pd.DataFrame()
    variant_map = aggregate_map_data(variants_df, "varianten") if variants_df is not None and not variants_df.empty else pd.DataFrame()
    parts = [part for part in [study_map, variant_map] if not part.empty]
    if not parts:
        return pd.DataFrame()
    combined = pd.concat(parts, ignore_index=True, sort=False)
    records = []
    keys = ["INSTELLINGSNAAM ACTUEEL", "GEMEENTENAAM", "Latitude", "Longitude"]
    for (institution, city, latitude, longitude), group in combined.groupby(keys, dropna=False, sort=True):
        total = int(group["Aantal onderdelen"].sum())
        blocks = []
        for _, item in group.iterrows():
            hover = item["Hover"]
            if not blocks:
                blocks.append(hover)
            else:
                pos = hover.find("<b>Type:")
                blocks.append(hover[pos:] if pos >= 0 else hover)
        records.append({
            "INSTELLINGSNAAM ACTUEEL": clean_text(institution), "GEMEENTENAAM": clean_text(city),
            "Latitude": float(latitude), "Longitude": float(longitude), "Aantal onderdelen": total,
            "Aantal label": f"{total} onderdeel" if total == 1 else f"{total} onderdelen",
            "Marker label": str(total) if total > 1 else "", "Hover": "<br><br>".join(blocks),
        })
    return apply_coordinate_offsets(pd.DataFrame(records))

def select_education_data(opleidingen, selected_levels):
    level_filtered = filter_levels(opleidingen, selected_levels)
    discipline, discipline_df = choose_discipline(level_filtered)
    clusters = get_cluster_values(discipline_df)
    options = [ALL_CLUSTERS, *clusters]
    if st.session_state.get("selected_education_cluster") not in options:
        st.session_state.selected_education_cluster = ALL_CLUSTERS
    st.sidebar.markdown("### Cluster")
    cluster = st.sidebar.selectbox("Cluster", options, key="selected_education_cluster", label_visibility="collapsed")
    filtered = filter_cluster(discipline_df, cluster)
    level_text = ", ".join(selected_levels)
    title = f"Opleidingen binnen {level_text}, discipline {discipline}" + ("" if cluster == ALL_CLUSTERS else f" en cluster {cluster}")
    return filtered, discipline, cluster, title


def select_variant_data(opleidingen, varianten, selected_levels):
    level_opleidingen = filter_levels(opleidingen, selected_levels)
    level_varianten = filter_levels(varianten, selected_levels)
    discipline, discipline_opleidingen = choose_discipline(level_opleidingen)
    education_clusters = get_cluster_values(discipline_opleidingen)
    if not education_clusters:
        st.warning("Binnen deze selectie zijn geen opleidingsclusters gevonden.")
        st.stop()
    # Cruciale overgang: 'Alle clusters' bestaat hier niet. Bij omschakelen wordt het eerste cluster gekozen.
    if st.session_state.get("selected_education_cluster") not in education_clusters:
        st.session_state.selected_education_cluster = education_clusters[0]
    st.sidebar.markdown("### Cluster opleiding")
    education_cluster = st.sidebar.selectbox("Cluster opleiding", education_clusters,
                                              key="selected_education_cluster", label_visibility="collapsed")
    selected_opleidingen, related_variants = variants_for_education_cluster(
        discipline_opleidingen, level_varianten, education_cluster)
    if related_variants.empty:
        st.warning("Voor het gekozen opleidingscluster zijn geen gekoppelde varianten gevonden.")
        st.stop()
    variant_clusters = get_cluster_values(related_variants)
    variant_options = [ALL_VARIANTS, *variant_clusters]
    # 'Alle varianten' is standaard bij openen of wanneer de oude keuze niet meer geldig is.
    if st.session_state.get("selected_variant_cluster") not in variant_options:
        st.session_state.selected_variant_cluster = ALL_VARIANTS
    st.sidebar.markdown("### Cluster variant")
    variant_cluster = st.sidebar.selectbox("Cluster variant", variant_options,
                                            key="selected_variant_cluster", label_visibility="collapsed")
    if variant_cluster == ALL_VARIANTS:
        filtered = related_variants.copy()
        displayed_opleidingen = selected_opleidingen.copy()
    else:
        # Specifieke variantclusters worden tenantbreed binnen de gekozen
        # opleidingsniveaus gezocht, dus buiten discipline en studiecluster.
        filtered = filter_cluster(level_varianten, variant_cluster)
        displayed_opleidingen = studies_for_selected_variants(
            level_opleidingen,
            filtered,
        )
    title = (
        f"Variantencluster {variant_cluster}"
        if variant_cluster != ALL_VARIANTS
        else f"Varianten binnen {discipline} - {education_cluster}"
    )
    return filtered, displayed_opleidingen, discipline, education_cluster, variant_cluster, title


def main():
    st.set_page_config(page_title="Kaart kunstopleidingen", layout="wide")
    st.title("Locaties van opleidingen en varianten")
    try:
        opleidingen, varianten = read_workbook(Path(__file__).with_name(DEFAULT_WORKBOOK))
    except Exception as error:
        st.error("Het Excelbestand kon niet correct worden ingelezen.")
        st.exception(error)
        st.stop()

    st.sidebar.markdown("### Kies gegevensbron")
    sheet = st.sidebar.radio("Kies gegevensbron", ["opleidingen", "varianten"], horizontal=True,
                             key="selected_sheet", label_visibility="collapsed")
    st.sidebar.divider()
    selected_levels = show_level_filter()
    if not selected_levels:
        st.stop()
    st.sidebar.divider()

    if sheet == "opleidingen":
        filtered, discipline, cluster, title = select_education_data(opleidingen, selected_levels)
        if filtered.empty:
            st.warning("Voor deze selectie zijn geen opleidingen gevonden.")
            st.stop()
        level_varianten = filter_levels(varianten, selected_levels)
        matching_variants = (
            filter_cluster(level_varianten, cluster)
            if cluster != ALL_CLUSTERS
            else level_varianten.iloc[0:0].copy()
        )
        if not matching_variants.empty:
            st.sidebar.info(
                f"Studiecluster '{cluster}' bestaat ook als variantcluster. "
                "De gelijknamige varianten, ook onder andere studieclusters, "
                "zijn automatisch aan de kaart toegevoegd."
            )
        st.sidebar.divider()
        show_list(filtered, "OPLEIDINGSNAAM ACTUEEL", "Studies in studiecluster", "Geen studies gevonden.")
        if not matching_variants.empty:
            st.sidebar.divider()
            show_list(matching_variants, "Variant", "Varianten in gelijknamig variantcluster", "Geen varianten gevonden.")
        st.sidebar.divider()
        labour_studies = pd.concat(
            [filtered, studies_for_selected_variants(opleidingen, matching_variants)],
            ignore_index=True,
            sort=False,
        ).drop_duplicates()
        show_labour_panel(labour_studies)
        export_opleidingen = filtered.copy()
        export_varianten = pd.concat(
            [variants_for_selected_studies(export_opleidingen, varianten), matching_variants],
            ignore_index=True, sort=False,
        ).drop_duplicates()
        filename = safe_filename(
            f"selectie_opleidingen_en_varianten_{discipline}_{cluster}.xlsx"
        )
        mode, record_label = "opleidingen", "opleidingsregel" if len(filtered) == 1 else "opleidingsregels"
        map_studies, map_variants = filtered, matching_variants
    else:
        filtered, selected_studies, discipline, education_cluster, variant_cluster, title = select_variant_data(
            opleidingen, varianten, selected_levels)
        if filtered.empty:
            st.warning("Voor deze selectie zijn geen varianten gevonden.")
            st.stop()
        level_opleidingen = filter_levels(opleidingen, selected_levels)
        matching_studies = (
            filter_cluster(level_opleidingen, variant_cluster)
            if variant_cluster != ALL_VARIANTS
            else level_opleidingen.iloc[0:0].copy()
        )
        if not matching_studies.empty:
            st.sidebar.info(
                f"Variantcluster '{variant_cluster}' bestaat ook als studiecluster. "
                "De opleidingen uit dit gelijknamige studiecluster zijn "
                "automatisch aan de kaart toegevoegd."
            )
        st.sidebar.divider()
        show_list(selected_studies, "OPLEIDINGSNAAM ACTUEEL", "Studies van geselecteerde varianten", "Geen studies gevonden.")
        if not matching_studies.empty:
            st.sidebar.divider()
            show_list(matching_studies, "OPLEIDINGSNAAM ACTUEEL", "Studies in gelijknamig studiecluster", "Geen studies gevonden.")
        st.sidebar.divider()
        show_list(filtered, "Variant", "Varianten onderdeel van selectie", "Geen varianten gevonden.")
        st.sidebar.divider()
        labour_studies = pd.concat(
            [selected_studies, matching_studies, studies_for_selected_variants(opleidingen, filtered)],
            ignore_index=True,
            sort=False,
        ).drop_duplicates()
        show_labour_panel(labour_studies)
        export_opleidingen = pd.concat(
            [selected_studies, matching_studies], ignore_index=True, sort=False,
        ).drop_duplicates()
        export_varianten = filtered.copy()
        filename = safe_filename(
            f"selectie_varianten_en_opleidingen_{discipline}_"
            f"{education_cluster}_{variant_cluster}.xlsx"
        )
        mode, record_label = "varianten", "variant" if len(filtered) == 1 else "varianten"
        map_studies, map_variants = matching_studies, filtered

    st.sidebar.divider()
    st.sidebar.download_button(
        "Download opleidingen en varianten als Excel",
        selection_to_excel(
            export_opleidingen,
            export_varianten,
        ),
        filename,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
    )
    map_data = combine_map_sources(map_studies, map_variants)
    if map_data.empty:
        st.warning("Voor deze selectie zijn geen locaties met coördinaten beschikbaar.")
        st.stop()
    st.plotly_chart(create_map_figure(map_data, title), use_container_width=True)
    locations = len(map_data)
    st.caption(f"{len(filtered)} {record_label} verdeeld over {locations} {'locatie' if locations == 1 else 'locaties'}. "
               "Een cijfer in een marker geeft aan dat meerdere opleidingen of varianten van dezelfde instelling "
               "op deze locatie binnen de huidige selectie vallen.")


if __name__ == "__main__":
    main()
